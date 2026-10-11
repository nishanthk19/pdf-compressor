"""Synthetic fixtures for offline extraction; no user documents or services."""
import json
import subprocess
import sys
import tempfile
from pathlib import Path
import pymupdf
from docx import Document

script = Path(__file__).resolve().parents[1] / 'scripts' / 'letter-extract.py'
with tempfile.TemporaryDirectory() as folder:
    root = Path(folder)
    text = 'Reference: QA/42\nDate: 10/10/2026\nDear Test Recipient\nPlease replace the sample scanner.\nTest Signatory'
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((60, 70), text, fontsize=16)
    pdf.save(root / 'searchable.pdf')
    scan = pymupdf.open()
    target = scan.new_page()
    target.insert_image(target.rect, pixmap=page.get_pixmap(matrix=pymupdf.Matrix(2, 2)))
    scan.save(root / 'scanned.pdf')
    docx = Document()
    for line in text.splitlines():
        docx.add_paragraph(line)
    table = docx.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Scanner'
    table.cell(0, 1).text = 'Quantity 1'
    docx.save(root / 'letter.docx')
    for name in ['searchable.pdf', 'scanned.pdf', 'letter.docx']:
        result = subprocess.run([sys.executable, str(script), str(root / name)], capture_output=True, text=True, timeout=90, check=True)
        data = json.loads(result.stdout)
        assert 'sample scanner' in data['text'], (name, data)
        assert 'QA/42' in data['text'], (name, data)
        assert data['metadata']['reviewRequired'] is True
        if name.endswith('.docx'):
            assert 'Scanner | Quantity 1' in data['text']
        print(name, 'extraction passed', flush=True)
